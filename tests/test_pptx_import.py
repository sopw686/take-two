"""PowerPoint notes import: decks built with python-pptx (dev only), read back with the stdlib parser."""

from __future__ import annotations

import io
import struct
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from take_two import pptx_import
from take_two.marks import parse_script
from take_two.pptx_import import PptxError, neutralize_line, notes_to_script

TITLE_SLIDE, TITLE_AND_CONTENT, TITLE_ONLY, BLANK = 0, 1, 5, 6  # python-pptx default template layouts


def deck(tmp_path, slides, name="deck.pptx") -> bytes:
    """slides: (layout, title or None, notes or None). notes=None leaves the slide without a notes slide."""
    prs = Presentation()
    for layout, title, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[layout])
        if title is not None:
            slide.shapes.title.text = title
        if notes is not None:
            slide.notes_slide.notes_text_frame.text = notes
    path = tmp_path / name
    prs.save(path)
    return path.read_bytes()


def rezip(data: bytes, replace: dict[str, bytes | None]) -> bytes:
    """Copy a zip, replacing (bytes) or dropping (None) the named members."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename not in replace:
                dst.writestr(info.filename, src.read(info))
            elif replace[info.filename] is not None:
                dst.writestr(info.filename, replace[info.filename])
    return out.getvalue()


def zip_of(members: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return out.getvalue()


def lines_by_section(text: str) -> list[list[str]]:
    s = parse_script(text)
    return [[ln.text for ln in s.section_lines(sec)] for sec in s.sections]


# ---- the script ---------------------------------------------------------------------

def test_one_section_per_slide_in_order(tmp_path):
    data = deck(tmp_path, [
        (TITLE_SLIDE, "Why reefs turn white", "Good morning.\nToday: what bleaching does to a reef."),
        (TITLE_AND_CONTENT, "Methods", None),
        (BLANK, None, "We sampled forty colonies."),
        (TITLE_ONLY, "Results", "Bleaching doubled.\n\n   \n  That is   the headline.  "),
    ])
    text, slides = notes_to_script(data)
    assert slides == 4
    assert text == ("## Slide 1: Why reefs turn white\nGood morning.\nToday: what bleaching does to a reef.\n\n"
                    "## Slide 2: Methods\n<!-- no notes -->\n\n"
                    "## Slide 3\nWe sampled forty colonies.\n\n"
                    "## Slide 4: Results\nBleaching doubled.\nThat is the headline.\n")


def test_the_mark_parser_reads_slides_as_sections_and_notes_as_plain_lines(tmp_path):
    data = deck(tmp_path, [
        (TITLE_SLIDE, "Coral bleaching", "Good morning.\nHeat / light both matter.\n## Key result\nCompare 1 // 2."),
        (TITLE_AND_CONTENT, "Methods", None),
        (BLANK, None, "One more thing."),
    ])
    text, _ = notes_to_script(data)
    s = parse_script(text)
    assert [sec.name for sec in s.sections] == ["Slide 1: Coral bleaching", "Slide 2: Methods", "Slide 3"]
    assert [sec.budget_s for sec in s.sections] == [None, None, None]
    assert lines_by_section(text) == [["Good morning.", "Heat/light both matter.", "Key result", "Compare 1//2."],
                                      [], ["One more thing."]]
    assert sum(len(ln.pauses) for ln in s.lines) == 0
    assert not any(ln.is_key for ln in s.lines) and not s.defines


def test_empty_notes_slide_and_empty_title_placeholder(tmp_path):
    data = deck(tmp_path, [(TITLE_ONLY, None, ""), (TITLE_ONLY, "  ", "\n\n")])
    text, slides = notes_to_script(data)
    assert slides == 2
    assert text == "## Slide 1\n<!-- no notes -->\n\n## Slide 2\n<!-- no notes -->\n"


def test_title_comes_from_the_title_placeholder_only(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = "Methods\vand data"  # \v is a line break (a:br) inside the paragraph
    slide.placeholders[1].text = "A bullet, not the title"
    slide.notes_slide.notes_text_frame.text = "first half\vsecond half"
    buf = io.BytesIO()
    prs.save(buf)
    text, _ = notes_to_script(buf.getvalue())
    assert text == "## Slide 1: Methods and data\nfirst half second half\n"


def test_slide_order_follows_the_presentation_not_part_names(tmp_path):
    prs = Presentation()
    for title in ("Made first", "Made second", "Made third"):
        prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY]).shapes.title.text = title
    sld_ids = prs.slides._sldIdLst
    first = sld_ids[0]
    sld_ids.remove(first)
    sld_ids.append(first)  # moved to the end, as dragging it in PowerPoint would
    path = tmp_path / "moved.pptx"
    prs.save(path)
    data = path.read_bytes()
    assert b"Made first" in zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml")
    text, _ = notes_to_script(data)
    assert [sec.name for sec in parse_script(text).sections] == [
        "Slide 1: Made second", "Slide 2: Made third", "Slide 3: Made first"]


def test_hash_heading_in_notes_stays_a_line():
    assert neutralize_line("## Key result") == "Key result"
    assert neutralize_line("  ### Also # this") == "Also # this"
    assert neutralize_line("##") == ""
    assert neutralize_line("# of samples: 40") == "# of samples: 40"  # one # never starts a section
    text = "## Slide 1\n" + neutralize_line("## Not a section") + "\n"
    assert lines_by_section(text) == [["Not a section"]]


def test_standalone_slashes_are_not_pauses():
    assert neutralize_line("heat / light") == "heat/light"
    assert neutralize_line("compare 1 // 2 and A/B") == "compare 1//2 and A/B"
    assert neutralize_line("/ at the start") == "/at the start"
    assert neutralize_line("at the end //") == "at the end//"
    assert neutralize_line("a / / b") == "a//b"
    assert neutralize_line("/ //") == ""
    s = parse_script("## Slide 1\n" + neutralize_line("Heat / light // matter /") + "\n")
    assert [ln.text for ln in s.lines] == ["Heat/light//matter/"]
    assert s.lines[0].pauses == []
    assert [n for t in s.lines[0].tokens for n in t.norm] == ["heat", "light", "matter"]


def test_comment_opener_in_notes_cannot_swallow_later_slides(tmp_path):
    # Unchanged, "<!--" would pair with the "-->" of slide 2's placeholder and blank slide 2's header.
    data = deck(tmp_path, [(TITLE_ONLY, "Intro", "<!-- remember to smile\nNext point."),
                           (TITLE_ONLY, "Methods", None),
                           (TITLE_ONLY, "Results <!-- draft", "Third.")])
    text, _ = notes_to_script(data)
    s = parse_script(text)
    assert [sec.name for sec in s.sections] == ["Slide 1: Intro", "Slide 2: Methods", "Slide 3: Results < !-- draft"]
    assert lines_by_section(text) == [["< !-- remember to smile", "Next point."], [], ["Third."]]


def test_title_ending_in_a_time_is_not_a_budget(tmp_path):
    text, _ = notes_to_script(deck(tmp_path, [(TITLE_ONLY, "Timeline [1:30]", "When it happened.")]))
    assert text.startswith("## Slide 1: Timeline (1:30)\n")
    assert parse_script(text).sections[0].budget_s is None


def test_hidden_slide_is_flagged_in_a_comment(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    slide.shapes.title.text = "Backup"
    slide.notes_slide.notes_text_frame.text = "Only if asked."
    slide._element.set("show", "0")
    buf = io.BytesIO()
    prs.save(buf)
    text, _ = notes_to_script(buf.getvalue())
    assert text == "## Slide 1: Backup\n<!-- hidden slide -->\nOnly if asked.\n"
    assert lines_by_section(text) == [["Only if asked."]]


def test_strict_open_xml_namespaces(tmp_path):
    data = deck(tmp_path, [(TITLE_ONLY, "Strict", "Saved as Strict Open XML.")])
    swaps = [(b"http://schemas.openxmlformats.org/presentationml/2006/main", b"http://purl.oclc.org/ooxml/presentationml/main"),
             (b"http://schemas.openxmlformats.org/drawingml/2006/main", b"http://purl.oclc.org/ooxml/drawingml/main"),
             (b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
              b"http://purl.oclc.org/ooxml/officeDocument/relationships")]
    src = zipfile.ZipFile(io.BytesIO(data))
    replace = {}
    for name in src.namelist():
        if name.endswith((".xml", ".rels")):
            raw = src.read(name)
            for a, b in swaps:
                raw = raw.replace(a, b)
            replace[name] = raw
    text, _ = notes_to_script(rezip(data, replace))
    assert text == "## Slide 1: Strict\nSaved as Strict Open XML.\n"


# ---- files that cannot be imported --------------------------------------------------

def raises(data: bytes, fragment: str) -> None:
    with pytest.raises(PptxError) as info:
        notes_to_script(data)
    assert fragment in info.value.message


def test_not_a_pptx():
    raises(b"", "empty")
    raises(b"definitely not a zip file " * 20, "not a PowerPoint (.pptx) file")
    raises(pptx_import.OLE_MAGIC + bytes(504), "password-protected or is an old .ppt")


def test_zip_without_a_presentation():
    raises(zip_of({}), "no ppt/presentation.xml")
    raises(zip_of({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"<w:document/>"}), "no ppt/presentation.xml")


def test_presentation_without_slides(tmp_path):
    raises(deck(tmp_path, []), "has no slides")


def test_oversized_part_is_refused_before_it_is_read(monkeypatch):
    big = zip_of({"ppt/presentation.xml": b" " * (pptx_import.MAX_PART_BYTES + 1)})  # 5 MB that deflates to a few KB
    assert len(big) < 100_000
    real_open = zipfile.ZipFile.open

    def guarded(self, name, *args, **kwargs):
        assert getattr(name, "filename", name) != "ppt/presentation.xml", "read before the size check"
        return real_open(self, name, *args, **kwargs)
    monkeypatch.setattr(zipfile.ZipFile, "open", guarded)
    raises(big, "larger than 5 MB once uncompressed")


def test_too_many_members():
    raises(zip_of({f"junk/{i}.xml": b"" for i in range(pptx_import.MAX_MEMBERS + 1)}), "more than 2000 parts")


def test_upload_and_total_xml_limits(tmp_path, monkeypatch):
    data = deck(tmp_path, [(TITLE_ONLY, "Fine", "Small notes.")])
    monkeypatch.setattr(pptx_import, "MAX_TOTAL_XML_BYTES", 5000)
    raises(data, "adds up to more than")
    monkeypatch.setattr(pptx_import, "MAX_UPLOAD_BYTES", len(data) - 1)
    raises(data, "import limit")


def test_doctype_in_a_part_is_refused(tmp_path):
    data = deck(tmp_path, [(TITLE_ONLY, "Fine", "Small notes.")])
    slide = zipfile.ZipFile(io.BytesIO(data)).read("ppt/slides/slide1.xml")
    head, rest = slide.split(b"?>", 1)
    evil = head + b'?><!DOCTYPE p:sld [<!ENTITY lol "lol">]>' + rest.replace(b"Fine", b"&lol;")
    raises(rezip(data, {"ppt/slides/slide1.xml": evil}), "DOCTYPE")


def test_damaged_parts(tmp_path):
    data = deck(tmp_path, [(TITLE_ONLY, "Fine", "Small notes.")])
    raises(rezip(data, {"ppt/slides/slide1.xml": b"<p:sld"}), "ppt/slides/slide1.xml is not valid XML")
    raises(rezip(data, {"ppt/slides/slide1.xml": None}), "slide 1 is listed but missing")
    # A notes slide that is referenced but missing is just a slide without notes.
    text, _ = notes_to_script(rezip(data, {"ppt/notesSlides/notesSlide1.xml": None}))
    assert text == "## Slide 1: Fine\n<!-- no notes -->\n"


def test_zip_encrypted_member(tmp_path):
    data = bytearray(deck(tmp_path, [(TITLE_ONLY, "Fine", "Small notes.")]))
    # Set the "encrypted" flag on presentation.xml's central directory entry (zipfile cannot write encrypted zips).
    pos = data.find(b"PK\x01\x02")
    while pos != -1:
        name_len = struct.unpack_from("<H", data, pos + 28)[0]
        if data[pos + 46:pos + 46 + name_len] == b"ppt/presentation.xml":
            data[pos + 8] |= 0x1
        pos = data.find(b"PK\x01\x02", pos + 1)
    raises(bytes(data), "password-protected")


# ---- route --------------------------------------------------------------------------

@pytest.fixture
def client():
    from take_two.app import app  # the real app, so the router is known to be mounted (before the static files)
    return TestClient(app, base_url="http://127.0.0.1:8765")


PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def test_route_returns_the_script(client, tmp_path):
    data = deck(tmp_path, [(TITLE_SLIDE, "Hello", "First line.\nSecond line."), (BLANK, None, None)])
    r = client.post("/api/import/pptx", files={"file": ("talk.pptx", data, PPTX_TYPE)})
    assert r.status_code == 200
    assert r.json() == {"text": "## Slide 1: Hello\nFirst line.\nSecond line.\n\n## Slide 2\n<!-- no notes -->\n",
                        "slides": 2}


def test_route_rejects_a_bad_file(client):
    r = client.post("/api/import/pptx", files={"file": ("talk.pptx", b"not a zip", PPTX_TYPE)})
    assert r.status_code == 400
    assert r.json()["detail"] == "This is not a PowerPoint (.pptx) file."


def test_route_rejects_an_oversized_upload(client, tmp_path, monkeypatch):
    data = deck(tmp_path, [(TITLE_ONLY, "Fine", "Small notes.")])
    monkeypatch.setattr(pptx_import, "MAX_UPLOAD_BYTES", 1000)
    r = client.post("/api/import/pptx", files={"file": ("talk.pptx", data, PPTX_TYPE)})
    assert r.status_code == 400
    assert "import limit" in r.json()["detail"]
